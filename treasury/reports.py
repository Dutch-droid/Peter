"""Report builders. Each returns a Report (title, columns, rows) so HTML, Excel and PDF share one result."""
from __future__ import annotations

import io
from dataclasses import dataclass, field
from datetime import date, timedelta
from decimal import Decimal

from sqlalchemy import select
from sqlalchemy.orm import Session

from . import config
from . import interest as I
from . import services as S
from .models import Account, Investor, InternetAccount, Lease, ProcurementRequest

D = Decimal


@dataclass
class Report:
    title: str
    columns: list[str]
    rows: list[list]
    subtitle: str = ""
    totals: list | None = None
    money_cols: set[int] = field(default_factory=set)


# param spec: (name, label, kind, default)  kinds: date, int, investor, account, quarter, optinvestor
def _today():
    return date.today()


def _q_default():
    t = _today()
    return (t.month - 1) // 3 + 1


REPORTS: dict[str, dict] = {}


def report(key, title, module, params):
    def deco(fn):
        REPORTS[key] = {"title": title, "module": module, "params": params, "fn": fn}
        return fn
    return deco


def _accounts(db: Session, investor_id: int | None = None):
    q = select(Account).order_by(Account.investor_id, Account.id)
    if investor_id:
        q = q.where(Account.investor_id == investor_id)
    return list(db.scalars(q))


def _sum(rows, idx):
    return sum((D(r[idx]) for r in rows), D(0))


def _fmt_status(a, as_of):
    return a.status(as_of)


@report("investor_statement", "Investor Statement", "treasury",
        [("investor_id", "Investor", "investor", None), ("as_of", "As of", "date", "today")])
def investor_statement(db, p):
    inv = db.get(Investor, int(p["investor_id"]))
    as_of = p["as_of"]
    rows = []
    for a in inv.accounts:
        t = a.current_term
        rows.append([a.reference, t.start_date, t.maturity_date, a.status(as_of), D(t.principal), D(t.rate_pct), S.accrued(a, as_of),
                     D(t.principal) + S.accrued(a, as_of) if not a.redemption_date else D(0)])
    tot = ["Total", "", "", "", _sum(rows, 4), "", _sum(rows, 6), _sum(rows, 7)]
    return Report(f"Investor Statement - {inv.name}", ["Account", "Start", "Maturity", "Status", "Principal", "Rate %", "Accrued interest", "Net position"],
                  rows, f"KRA PIN {inv.kra_pin} | Bank {inv.bank_name} {inv.bank_account} | as of {as_of}", tot, {4, 6, 7})


@report("investment_statement", "Investment Statement (per account)", "treasury",
        [("account_id", "Account", "account", None), ("as_of", "As of", "date", "today")])
def investment_statement(db, p):
    a = db.get(Account, int(p["account_id"]))
    as_of = p["as_of"]
    rows = []
    for t in a.terms:
        term = I.Term(D(t.principal), D(t.rate_pct), t.start_date, t.maturity_date)
        rows.append([f"Term {t.seq}", t.start_date, t.maturity_date, D(t.principal), D(t.rate_pct),
                     I.period_interest([term], t.start_date, t.maturity_date, a.redemption_date, as_of)])
    for adj in a.adjustments:
        rows.append([f"Adjustment: {adj.reason}", adj.effective_date, "", "", "", D(adj.amount)])
    return Report(f"Investment Statement - {a.reference} ({a.investor.name})", ["Item", "From", "To", "Principal", "Rate %", "Interest"],
                  rows, f"Status {a.status(as_of)} | as of {as_of}", ["Total accrued", "", "", "", "", S.accrued(a, as_of)], {3, 5})


@report("interest_earned", "Interest Earned Report", "treasury",
        [("start", "From", "date", "month_start"), ("end", "To (inclusive)", "date", "today")])
def interest_earned(db, p):
    s, e = p["start"], p["end"] + timedelta(days=1)
    rows = []
    for a in _accounts(db):
        g = S.period_gross(a, s, e)
        if g:
            rows.append([a.investor.name, a.reference, D(a.current_term.principal), D(a.current_term.rate_pct), g])
    return Report("Interest Earned", ["Investor", "Account", "Principal", "Rate %", "Interest earned"], rows,
                  f"{p['start']} to {p['end']}", ["Total", "", _sum(rows, 2), "", _sum(rows, 4)], {2, 4})


@report("active_investors", "Active Investors", "treasury", [("as_of", "As of", "date", "today")])
def active_investors(db, p):
    as_of, rows = p["as_of"], []
    for inv in db.scalars(select(Investor).order_by(Investor.name)):
        act = [a for a in inv.accounts if a.status(as_of) in ("Active", "Matured", "Pending")]
        if act:
            rows.append([inv.name, inv.kra_pin, inv.telephone, inv.email, len(act), sum((S.principal_outstanding(a, as_of) for a in act), D(0))])
    return Report("Active Investors", ["Investor", "KRA PIN", "Telephone", "Email", "Accounts", "Principal"], rows, f"as of {as_of}",
                  ["Total", "", "", "", sum(r[4] for r in rows), _sum(rows, 5)], {5})


@report("total_funds", "Total Funds Invested", "treasury", [("as_of", "As of", "date", "today")])
def total_funds(db, p):
    as_of, rows = p["as_of"], []
    for inv in db.scalars(select(Investor).order_by(Investor.name)):
        amt = sum((S.principal_outstanding(a, as_of) for a in inv.accounts), D(0))
        if amt:
            rows.append([inv.name, sum(1 for a in inv.accounts if S.principal_outstanding(a, as_of) > 0), amt])
    return Report("Total Funds Invested", ["Investor", "Open accounts", "Principal"], rows, f"as of {as_of}",
                  ["Total", sum(r[1] for r in rows), _sum(rows, 2)], {2})


@report("accrual_summary", "Interest Accrual Summary", "treasury", [("as_of", "As of", "date", "today")])
def accrual_summary(db, p):
    as_of = p["as_of"]
    ms, _ = I.month_bounds(as_of.year, as_of.month)
    qs, _ = I.quarter_bounds(as_of.year, (as_of.month - 1) // 3 + 1)
    rows = []
    for inv in db.scalars(select(Investor).order_by(Investor.name)):
        accts = inv.accounts
        if not accts:
            continue
        day = sum((S.period_gross(a, as_of - timedelta(days=1), as_of, as_of) for a in accts), D(0))
        rows.append([inv.name, len(accts), day, sum((S.period_gross(a, ms, as_of, as_of) for a in accts), D(0)),
                     sum((S.period_gross(a, qs, as_of, as_of) for a in accts), D(0)), sum((S.accrued(a, as_of) for a in accts), D(0))])
    return Report("Interest Accrual Summary", ["Investor", "Accounts", "Last day", "Month to date", "Quarter to date", "Cumulative"], rows,
                  f"as of {as_of} (accrual through the day before)", ["Total", sum(r[1] for r in rows)] + [_sum(rows, i) for i in (2, 3, 4, 5)], {2, 3, 4, 5})


def _quarterly(db, p):
    y, q = int(p["year"]), int(p["quarter"])
    s, e = I.quarter_bounds(y, q)
    inv_id = int(p["investor_id"]) if p.get("investor_id") else None
    rows = []
    for a in _accounts(db, inv_id):
        g = S.period_gross(a, s, e, as_of=_today() if e > _today() else None)
        if g == 0 and not a.terms:
            continue
        wht = I.withholding_tax(g)
        rows.append([a.investor.name, a.reference, D(a.current_term.principal), D(a.current_term.rate_pct), g, wht, g - wht])
    if p.get("group") == "investor":
        agg: dict[str, list] = {}
        for r in rows:
            x = agg.setdefault(r[0], [r[0], "(all accounts)", D(0), "", D(0), D(0), D(0)])
            x[2] += r[2]; x[4] += r[4]; x[5] += r[5]; x[6] += r[6]
        rows = list(agg.values())
    rows = [r for r in rows if r[4] != 0]
    return y, q, s, e, rows


@report("quarterly_interest", "Quarterly Interest Report", "treasury",
        [("year", "Year", "int", "year"), ("quarter", "Quarter", "quarter", None), ("investor_id", "Investor (optional)", "optinvestor", None),
         ("group", "Group by", "group", "account")])
def quarterly_interest(db, p):
    y, q, s, e, rows = _quarterly(db, p)
    return Report(f"Quarterly Interest Report - Q{q} {y}", ["Investor", "Account", "Principal", "Rate %", "Gross interest", f"WHT ({config.WHT_RATE:.0%})", "Net interest payable"],
                  rows, f"{s} to {e - timedelta(days=1)}", ["Total", "", "", "", _sum(rows, 4), _sum(rows, 5), _sum(rows, 6)], {2, 4, 5, 6})


@report("withholding_tax", "Withholding Tax Report", "treasury",
        [("year", "Year", "int", "year"), ("quarter", "Quarter", "quarter", None)])
def withholding_tax_report(db, p):
    y, q, s, e, rows = _quarterly(db, {**p, "group": "investor"})
    inv = {i.name: i.kra_pin for i in db.scalars(select(Investor))}
    out = [[r[0], inv.get(r[0], ""), r[4], r[5]] for r in rows]
    return Report(f"Withholding Tax Report - Q{q} {y}", ["Investor", "KRA PIN", "Gross interest", "Withholding tax"], out,
                  f"{s} to {e - timedelta(days=1)}", ["Total", "", _sum(out, 2), _sum(out, 3)], {2, 3})


@report("net_position", "Net Position Report", "treasury", [("as_of", "As of", "date", "today")])
def net_position(db, p):
    as_of, rows = p["as_of"], []
    for inv in db.scalars(select(Investor).order_by(Investor.name)):
        live = [a for a in inv.accounts if not a.redemption_date]
        pr = sum((S.principal_outstanding(a, as_of) for a in live), D(0))
        ia = sum((S.accrued(a, as_of) for a in live), D(0))
        if live:
            rows.append([inv.name, pr, ia, pr + ia])
    return Report("Net Position Report", ["Investor", "Total principal", "Total accrued interest", "Net position"], rows,
                  f"as of {as_of} | Net position = principal + accrued interest (non-redeemed accounts)",
                  ["Total", _sum(rows, 1), _sum(rows, 2), _sum(rows, 3)], {1, 2, 3})


@report("maturing", "Maturing Investments", "treasury", [("days", "Maturing within (days)", "int", 30)])
def maturing(db, p):
    today, lim, rows = _today(), _today() + timedelta(days=int(p["days"])), []
    for a in _accounts(db):
        if a.redemption_date:
            continue
        t = a.current_term
        if t.maturity_date <= lim:
            rows.append([a.investor.name, a.reference, D(t.principal), D(t.rate_pct), t.maturity_date, (t.maturity_date - today).days, S.accrued(a, today)])
    rows.sort(key=lambda r: r[4])
    return Report("Maturing Investments", ["Investor", "Account", "Principal", "Rate %", "Maturity", "Days left", "Accrued interest"], rows,
                  f"maturing on or before {lim} (negative days = already matured, action required)", None, {2, 6})


@report("redeemed", "Redeemed Investments", "treasury", [("start", "From", "date", "year_start"), ("end", "To", "date", "today")])
def redeemed(db, p):
    rows = []
    for a in _accounts(db):
        if a.redemption_date and p["start"] <= a.redemption_date <= p["end"]:
            rows.append([a.investor.name, a.reference, D(a.current_term.principal), a.redemption_date, S.accrued(a, a.redemption_date)])
    return Report("Redeemed Investments", ["Investor", "Account", "Principal", "Redeemed on", "Total interest earned"], rows,
                  f"{p['start']} to {p['end']}", ["Total", "", _sum(rows, 2), "", _sum(rows, 4)], {2, 4})


@report("vat_report", "VAT & Procurement Report", "finance", [("start", "From", "date", "year_start"), ("end", "To", "date", "today")])
def vat_report(db, p):
    q = select(ProcurementRequest).where(ProcurementRequest.request_date >= p["start"], ProcurementRequest.request_date <= p["end"]).order_by(ProcurementRequest.request_date)
    rows = [[r.request_date, r.department, r.supplier.name, r.approval_status, D(r.total_cost), D(r.vat_amount), D(r.invoice_value)] for r in db.scalars(q)]
    return Report("VAT & Procurement Report", ["Date", "Department", "Supplier", "Status", "Net cost", "VAT", "Invoice value"], rows,
                  f"{p['start']} to {p['end']}", ["Total", "", "", "", _sum(rows, 4), _sum(rows, 5), _sum(rows, 6)], {4, 5, 6})


@report("rent_schedule", "Rent Schedule (next 12 months)", "finance", [])
def rent_schedule(db, p):
    rows, end = [], _today() + timedelta(days=366)
    for l in db.scalars(select(Lease)):
        d = l.next_due_date
        while d <= end and (l.lease_expiry is None or d <= l.lease_expiry):
            rows.append([d, l.branch_name, l.landlord, l.frequency, D(l.rent_amount), l.pay_point])
            d = I.add_period(d, l.frequency)
    rows.sort(key=lambda r: r[0])
    return Report("Rent Schedule", ["Due date", "Branch", "Landlord", "Frequency", "Amount", "Pay point"], rows, "projected from next due date until lease expiry",
                  None, {4})


@report("internet_schedule", "Internet Billing Schedule (next 12 months)", "finance", [])
def internet_schedule(db, p):
    rows, end = [], _today() + timedelta(days=366)
    for a in db.scalars(select(InternetAccount)):
        d = a.next_due_date
        while d <= end:
            rows.append([d, a.branch_name, a.isp, a.account_number, D(a.monthly_cost), a.pay_point])
            d = I.add_period(d, "Monthly")
    rows.sort(key=lambda r: r[0])
    return Report("Internet Billing Schedule", ["Due date", "Branch", "ISP", "Account", "Amount", "Pay point"], rows, "projected monthly", None, {4})


def parse_params(spec, raw) -> dict:
    t = _today()
    defaults = {"today": t, "month_start": t.replace(day=1), "year_start": date(t.year, 1, 1), "year": t.year}
    out = {}
    for name, _label, kind, default in spec:
        v = raw.get(name)
        if kind == "date":
            out[name] = date.fromisoformat(v) if v else defaults.get(default, t)
        elif kind in ("int", "quarter"):
            out[name] = int(v) if v else (defaults.get(default) if default in defaults else (_q_default() if kind == "quarter" else default))
        elif kind == "group":
            out[name] = v or default
        else:
            out[name] = v or default
    return out


def to_xlsx(r: Report) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = r.title[:31]
    ws.append([r.title]); ws["A1"].font = Font(bold=True, size=13)
    ws.append([r.subtitle])
    ws.append([])
    ws.append(r.columns)
    for c in ws[4]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="1F3A5F")
    for row in r.rows + ([r.totals] if r.totals else []):
        ws.append([float(v) if isinstance(v, Decimal) else v for v in row])
    for row in ws.iter_rows(min_row=5):
        for i, c in enumerate(row):
            if i in r.money_cols and isinstance(c.value, (int, float)):
                c.number_format = "#,##0.00"
    if r.totals:
        for c in ws[ws.max_row]:
            c.font = Font(bold=True)
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = min(48, max(len(str(c.value or "")) for c in col[3:]) + 2) if ws.max_row > 3 else 14
    buf = io.BytesIO(); wb.save(buf)
    return buf.getvalue()


def to_pdf(r: Report) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    ss = getSampleStyleSheet()
    cell = ParagraphStyle("c", parent=ss["BodyText"], fontSize=7.5, leading=9)
    hdr = ParagraphStyle("h", parent=cell, textColor=colors.white, fontName="Helvetica-Bold")
    def fmt(i, v):
        if isinstance(v, Decimal):
            return f"{v:,.2f}"
        return "" if v is None else str(v)
    data = [[Paragraph(c, hdr) for c in r.columns]]
    for row in r.rows + ([r.totals] if r.totals else []):
        data.append([Paragraph(fmt(i, v), cell) for i, v in enumerate(row)])
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=14 * mm, rightMargin=14 * mm, topMargin=14 * mm, bottomMargin=14 * mm, title=r.title)
    t = Table(data, repeatRows=1)
    st = [("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F3A5F")), ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#BBBBBB")),
          ("VALIGN", (0, 0), (-1, -1), "TOP")]
    if r.totals:
        st.append(("BACKGROUND", (0, len(data) - 1), (-1, len(data) - 1), colors.HexColor("#E8EEF5")))
    t.setStyle(TableStyle(st))
    doc.build([Paragraph(r.title, ss["Title"]), Paragraph(r.subtitle, ss["BodyText"]), Spacer(1, 6), t,
               Spacer(1, 8), Paragraph(f"Generated {date.today()}", cell)])
    return buf.getvalue()
